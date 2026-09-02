#!/usr/bin/env python3
"""Build a new .docx from a chunk plan (used for PDF sources, which have no
existing docx structure to edit in place).

Usage:
    python3 build_docx_from_plan.py <plan.json> <output.docx>

plan.json:
    {
      "chunk_groups": [
        {"elements": [
          {"type": "paragraph", "style": "Heading 1", "text": "..."},
          {"type": "paragraph", "text": "..."},
          {"type": "table", "rows": [["a", "b"], ["c", "d"]]}
        ]},
        {"elements": [{"type": "paragraph", "text": "..."}]}
      ]
    }

Each chunk_groups entry becomes one final chunk. A standalone "<<<CHUNK>>>"
marker paragraph is inserted between consecutive groups (not before the
first, not after the last). Ends with a structural self-check.
"""

import json
import sys

ALLOWED_STYLES = {"Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "List Bullet"}


def load_plan(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        print(f"error: could not read plan '{path}': {e}", file=sys.stderr)
        sys.exit(1)


def validate_plan(plan):
    errors = []
    chunk_groups = plan.get("chunk_groups")
    if not isinstance(chunk_groups, list) or not chunk_groups:
        errors.append("'chunk_groups' must be a non-empty list")
        return errors, []

    total_tables = 0
    for gi, group in enumerate(chunk_groups):
        elements = group.get("elements") if isinstance(group, dict) else None
        if not isinstance(elements, list) or not elements:
            errors.append(f"chunk_groups[{gi}].elements must be a non-empty list")
            continue
        for ei, el in enumerate(elements):
            if not isinstance(el, dict) or "type" not in el:
                errors.append(f"chunk_groups[{gi}].elements[{ei}] missing 'type'")
                continue
            etype = el["type"]
            if etype == "paragraph":
                if not isinstance(el.get("text"), str):
                    errors.append(f"chunk_groups[{gi}].elements[{ei}] paragraph missing string 'text'")
            elif etype == "table":
                rows = el.get("rows")
                if not isinstance(rows, list) or not rows:
                    errors.append(f"chunk_groups[{gi}].elements[{ei}] table 'rows' must be a non-empty list")
                    continue
                width = None
                for ri, row in enumerate(rows):
                    if not isinstance(row, list) or not all(isinstance(c, str) for c in row):
                        errors.append(f"chunk_groups[{gi}].elements[{ei}].rows[{ri}] must be a list of strings")
                        continue
                    if width is None:
                        width = len(row)
                    elif len(row) != width:
                        errors.append(
                            f"chunk_groups[{gi}].elements[{ei}].rows[{ri}] has {len(row)} cols, "
                            f"expected {width} (jagged table)"
                        )
                total_tables += 1
            else:
                errors.append(f"chunk_groups[{gi}].elements[{ei}] has unknown type {etype!r}")

    return errors, chunk_groups


def main():
    if len(sys.argv) != 3:
        print("usage: build_docx_from_plan.py <plan.json> <output.docx>", file=sys.stderr)
        sys.exit(1)

    plan_path, output_path = sys.argv[1], sys.argv[2]

    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        print("error: python-docx is not installed (pip install python-docx)", file=sys.stderr)
        sys.exit(1)

    plan = load_plan(plan_path)
    errors, chunk_groups = validate_plan(plan)
    if errors:
        for e in errors:
            print(f"error: {e}", file=sys.stderr)
        sys.exit(1)

    doc = Document()
    plan_table_count = 0

    for gi, group in enumerate(chunk_groups):
        for el in group["elements"]:
            if el["type"] == "paragraph":
                style = el.get("style") or "Normal"
                if style not in ALLOWED_STYLES:
                    print(
                        f"warning: chunk_groups[{gi}] paragraph style {style!r} not in allowed set, "
                        "falling back to 'Normal'",
                        file=sys.stderr,
                    )
                    style = "Normal"
                doc.add_paragraph(el["text"], style=style)
            else:  # table
                rows = el["rows"]
                n_rows = len(rows)
                n_cols = len(rows[0])
                table = doc.add_table(rows=n_rows, cols=n_cols)
                for r, row in enumerate(rows):
                    for c, cell_text in enumerate(row):
                        table.cell(r, c).text = cell_text
                plan_table_count += 1

        if gi != len(chunk_groups) - 1:
            doc.add_paragraph("<<<CHUNK>>>")

    doc.save(output_path)

    # --- structural self-check ---
    self_check_failures = []
    try:
        out_doc = Document(output_path)
    except Exception as e:
        print(f"SELF-CHECK: FAIL - could not reopen output file: {e}", file=sys.stderr)
        sys.exit(1)

    out_items = list(out_doc.iter_inner_content())
    out_table_count = sum(1 for it in out_items if isinstance(it, Table))
    if out_table_count != plan_table_count:
        self_check_failures.append(f"expected {plan_table_count} tables in output, found {out_table_count}")

    marker_positions = [
        i for i, it in enumerate(out_items)
        if isinstance(it, Paragraph) and it.text.strip() == "<<<CHUNK>>>"
    ]
    expected_markers = len(chunk_groups) - 1
    if len(marker_positions) != expected_markers:
        self_check_failures.append(
            f"expected {expected_markers} marker paragraphs, found {len(marker_positions)}"
        )
    if marker_positions and marker_positions[0] == 0:
        self_check_failures.append("a marker paragraph is the first element in the document body")
    if marker_positions and marker_positions[-1] == len(out_items) - 1:
        self_check_failures.append("a marker paragraph is the last element in the document body")

    if self_check_failures:
        print("SELF-CHECK: FAIL")
        for f in self_check_failures:
            print(f"  - {f}", file=sys.stderr)
        sys.exit(1)

    print("SELF-CHECK: PASS")


if __name__ == "__main__":
    main()
